import openpyxl
import csv
import os

dir_path = '/Users/junichi.takayama/Desktop/案件管理/DBJ_社史Webサイト提案/03_Communication/見積もり/'

f1 = dir_path + '20260820_日本政策投資銀行様_スケジュール案 (2) (1).xlsx'
f2 = dir_path + '【見積／原価】2608_DBJ(日本政策投資銀行）_社史サイト - 260826_原価見積_すべてhtmlで制作した場合.csv'
f3 = dir_path + 'est_DBJ_0828.xlsx'

print("=== FILE 1 ===")
print("File:", os.path.basename(f1))
try:
    wb1 = openpyxl.load_workbook(f1, data_only=True)
    print("Sheets:", wb1.sheetnames)
    for sname in wb1.sheetnames:
        ws = wb1[sname]
        print(f"--- Sheet: {sname} ---")
        for r in range(1, ws.max_row+1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
            non_empty = [f"Col{c}:{v}" for c, v in enumerate(row_vals, 1) if v is not None]
            if non_empty:
                line = " | ".join(non_empty)
                if any(k in line for k in ["原価", "合計", "小計", "見積", "金額", "時間", "工数", "総計"]):
                    print(f"Row {r:3d}: {line}")
except Exception as e:
    print("Error loading f1:", e)

print("\n=== FILE 2 ===")
print("File:", os.path.basename(f2))
try:
    with open(f2, encoding='utf-8', errors='ignore') as f:
        reader = csv.reader(f)
        for r, row in enumerate(reader, 1):
            non_empty = [f"Col{c}:{v}" for c, v in enumerate(row, 1) if v != ""]
            if non_empty:
                line = " | ".join(non_empty)
                if any(k in line for k in ["原価", "合計", "小計", "見積", "金額", "時間", "工数", "総計"]):
                    print(f"Row {r:3d}: {line}")
except Exception as e:
    print("Error loading f2:", e)

print("\n=== FILE 3 ===")
print("File:", os.path.basename(f3))
try:
    wb3 = openpyxl.load_workbook(f3, data_only=True)
    print("Sheets:", wb3.sheetnames)
    for sname in wb3.sheetnames:
        ws = wb3[sname]
        print(f"--- Sheet: {sname} ---")
        for r in range(1, ws.max_row+1):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
            non_empty = [f"Col{c}:{v}" for c, v in enumerate(row_vals, 1) if v is not None]
            if non_empty:
                line = " | ".join(non_empty)
                if any(k in line for k in ["原価", "合計", "小計", "見積", "金額", "時間", "工数", "総計"]):
                    print(f"Row {r:3d}: {line}")
except Exception as e:
    print("Error loading f3:", e)
