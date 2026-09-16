import pptx
import sys

pptx_path = '/Users/junichi.takayama/Desktop/案件管理/DBJ_社史Webサイト提案/02_Materials/0826_オロ様ご確認【株式会社日本政策投資銀行御中】周年サイト提案書.pptx'

prs = pptx.Presentation(pptx_path)
print(f"Total Slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = ''.join(run.text for run in p.runs).strip()
                if t:
                    texts.append(t)
        elif shape.has_table:
            for row in shape.table.rows:
                r_text = ' | '.join(cell.text.strip().replace('\n', ' ') for cell in row.cells)
                if r_text:
                    texts.append(r_text)
    full_text = ' / '.join(texts)
    title = texts[0] if texts else '(Empty)'
    
    keywords = ['オロ', 'XXXX', 'xxxx', '黄色', '確認', '記載', '未定', 'ダミー']
    found = [k for k in keywords if k in full_text]
    flag_str = f" [★該当キーワード: {','.join(found)}]" if found else ""
    
    print(f"Slide {i+1:02d}: {title}{flag_str}")
