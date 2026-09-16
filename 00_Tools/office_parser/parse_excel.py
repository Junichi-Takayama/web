
import openpyxl
import json
import argparse

def excel_to_markdown(filepath):
    workbook = openpyxl.load_workbook(filepath)
    markdown_output = []

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        markdown_output.append(f"## Sheet: {sheet_name}\n")

        # ヘッダー行の取得
        header = [cell.value if cell.value is not None else "" for cell in sheet[1]]
        markdown_output.append("| " + " | ".join(map(str, header)) + " |")
        markdown_output.append("|" + "---|".join(["-" * len(str(h)) for h in header]) + "|")

        # データ行の取得
        for row in sheet.iter_rows(min_row=2):
            row_values = [cell.value if cell.value is not None else "" for cell in row]
            markdown_output.append("| " + " | ".join(map(str, row_values)) + " |")
        markdown_output.append("\n")
    return "\n".join(markdown_output)

def excel_to_json(filepath):
    workbook = openpyxl.load_workbook(filepath)
    json_output = {}

    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        data = []
        headers = [cell.value if cell.value is not None else "" for cell in sheet[1]]

        for row_index, row in enumerate(sheet.iter_rows(min_row=2)):
            row_data = {}
            for col_index, cell in enumerate(row):
                if col_index < len(headers):
                    row_data[headers[col_index]] = cell.value
            data.append(row_data)
        json_output[sheet_name] = data
    return json.dumps(json_output, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Parse Excel files and convert to Markdown or JSON.")
    parser.add_argument("filepath", help="Path to the Excel file (.xlsx).")
    parser.add_argument("--format", choices=["markdown", "json"], default="markdown",
                        help="Output format (markdown or json). Default is markdown.")
    args = parser.parse_args()

    if args.format == "markdown":
        print(excel_to_markdown(args.filepath))
    elif args.format == "json":
        print(excel_to_json(args.filepath))
